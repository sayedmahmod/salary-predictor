"""BA Entgeltstatistik (Jahreszahlen) -> aggregates.

Die Bundesagentur veroeffentlicht keine Quantile, sondern eine
Klassenbesetzung: wie viele sozialversicherungspflichtig Vollzeitbeschaeftigte
der Kerngruppe in jeder Entgeltklasse liegen, plus den Median.

Daraus werden p10/p25/p75/p90 durch lineare Interpolation innerhalb der
Klasse rekonstruiert (``quantile_method = interpolated_from_brackets``); der
Median wird direkt uebernommen. Faellt ein Quantil in die oben offene Klasse
"ueber 6.000 EUR", bleibt es leer - eine Zahl waere dort geraten.

Verwendet werden Blatt ``8.4`` (Region x KldB-2010-Beruf) und Blatt ``5.1``
(Deutschland x Beruf x Geschlecht). Bezugsgroesse ist das **Brutto-
monatsentgelt**; auf Jahresbasis wird mit Faktor 12 gerechnet.
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import pandas as pd

from salarykit import paths, schema
from salarykit.money import FxTable
from salarykit.sources._common import stable_id
from salarykit.titles.taxonomy import ROLES, role

SOURCE = "ba_entgelt"
DATASET = "Bundesagentur fuer Arbeit, Entgeltstatistik (Jahreszahlen)"
LICENSE = "dl-de/by-2-0 (Statistik der Bundesagentur fuer Arbeit)"

#: Untergrenzen der Entgeltklassen in EUR/Monat; die letzte Klasse ist offen.
BRACKET_EDGES = [0.0, 2000.0, 3000.0, 4000.0, 5000.0, 6000.0, None]
MONTHS_PER_YEAR = 12

_YEAR_IN_NAME = re.compile(r"-(\d{4})\d{2}-")
_KLDB_PREFIX = re.compile(r"^(?:dar\.\s*)?(\d{2,4})\s*(.*)$")


def quantile_from_brackets(counts: list[float], q: float) -> float | None:
    """Quantil aus einer Klassenbesetzung linear interpolieren.

    ``counts`` sind die Besetzungszahlen der Klassen in ``BRACKET_EDGES``.
    Liegt das Quantil in der oben offenen Klasse, wird ``None`` geliefert.
    """
    total = sum(c for c in counts if c)
    if total <= 0:
        return None
    target = q * total
    cumulative = 0.0
    for i, count in enumerate(counts):
        count = count or 0.0
        if cumulative + count < target:
            cumulative += count
            continue
        low = BRACKET_EDGES[i]
        high = BRACKET_EDGES[i + 1] if i + 1 < len(BRACKET_EDGES) else None
        if high is None or count <= 0:
            return None                      # offene Randklasse
        return low + (target - cumulative) / count * (high - low)
    return None


def _num(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(".", "").replace(",", ".")
    if not text or not re.match(r"^-?\d+(\.\d+)?$", text):
        return None
    return float(text)


def _kldb_lookup() -> dict[str, str]:
    out: dict[str, str] = {}
    for r in ROLES:
        if r.kldb2010:
            out.setdefault(r.kldb2010, r.code)
    return out


def _split_occupation(value) -> tuple[str | None, str | None]:
    """"27 Techn.Entwickl..." -> ("27", "Techn.Entwickl...")."""
    if value is None:
        return None, None
    text = str(value).strip()
    if not text:
        return None, None
    if text.lower().startswith("insgesamt"):
        return "TOTAL", "Insgesamt"
    hit = _KLDB_PREFIX.match(text)
    if hit:
        return hit.group(1), hit.group(2).strip() or text
    return None, text


def _rows_from_regional_sheet(sheet, year: int) -> list[dict]:
    rows: list[dict] = []
    for raw in sheet.iter_rows(min_row=10, values_only=True):
        if raw[0] is None or raw[1] is None:
            continue
        code, label = _split_occupation(raw[2])
        if code is None:
            continue
        counts = [_num(raw[i]) for i in range(4, 10)]
        rows.append({
            "geo_code": str(raw[0]).strip(), "geo_name": str(raw[1]).strip(),
            "occupation_code": code, "occupation_label": label,
            "employment_count": _num(raw[3]), "counts": counts,
            "median": _num(raw[10]), "sex": "total", "year": year,
        })
    return rows


def _rows_from_sex_sheet(sheet, year: int) -> list[dict]:
    """Blatt 5.1: drei Bloecke (Insgesamt / Maenner / Frauen) untereinander."""
    sex = "total"
    blocks = {"insgesamt": "total", "männer": "male", "frauen": "female"}
    rows: list[dict] = []
    for raw in sheet.iter_rows(min_row=11, values_only=True):
        first = str(raw[0]).strip().lower() if raw[0] is not None else ""
        if first in blocks:
            sex = blocks[first]
            continue
        code, label = _split_occupation(
            f"{raw[0]} {raw[1]}" if raw[0] is not None else raw[1])
        if code is None or code == "TOTAL":
            continue
        counts = [_num(raw[i]) for i in range(3, 9)]
        rows.append({
            "geo_code": "D", "geo_name": "Deutschland",
            "occupation_code": code, "occupation_label": label,
            "employment_count": _num(raw[2]), "counts": counts,
            "median": _num(raw[9]), "sex": sex, "year": year,
        })
    return rows


def _read_file(path: Path) -> list[dict]:
    hit = _YEAR_IN_NAME.search(path.name)
    if not hit:
        return []
    year = int(hit.group(1))
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows: list[dict] = []
    if "8.4" in workbook.sheetnames:
        rows += _rows_from_regional_sheet(workbook["8.4"], year)
    if "5.1" in workbook.sheetnames:
        rows += _rows_from_sex_sheet(workbook["5.1"], year)
    workbook.close()
    return rows


def load(fx: FxTable, predictor=None) -> pd.DataFrame:
    folder = paths.RAW_BA / "jahreszahlen"
    if not folder.exists():
        return pd.DataFrame()

    rows: list[dict] = []
    # Externe/macOS-Dateisysteme legen AppleDouble-Metadaten als
    # ``._datei.xlsx`` daneben. Das sind keine ZIP/XLSX-Arbeitsmappen.
    for path in sorted(p for p in folder.glob("*.xlsx")
                       if not p.name.startswith("._")):
        rows += _read_file(path)
    if not rows:
        return pd.DataFrame()

    df = pd.DataFrame(rows)
    # Blatt 5.1 (Deutschland gesamt) und 8.4 ueberschneiden sich beim
    # Geschlecht "total" - dort gewinnt 8.4, weil es die Regionen mitbringt.
    df = df.drop_duplicates(subset=["year", "geo_code", "occupation_code", "sex"],
                            keep="first").reset_index(drop=True)

    quantiles = {q: [quantile_from_brackets(c, q) for c in df.counts]
                 for q in (0.10, 0.25, 0.75, 0.90)}

    kldb = _kldb_lookup()
    codes = df.occupation_code.astype("string")
    mapped = codes.map(kldb)
    mapped = mapped.fillna(codes.str.slice(0, 3).map(kldb))
    mapped = mapped.fillna(codes.str.slice(0, 2).map(kldb))
    method = mapped.notna().map({True: "kldb_crosswalk", False: None})

    if predictor is not None and mapped.isna().any():
        todo = df.loc[mapped.isna(), "occupation_label"].astype("string").fillna("")
        guesses = {t: p for t, p in
                   zip(todo.drop_duplicates(),
                       predictor.predict_many(todo.drop_duplicates().tolist()))}
        for idx, label in todo.items():
            pred = guesses.get(label)
            if pred is not None and pred.method != "unmatched":
                mapped.at[idx] = pred.code
                method.at[idx] = pred.method

    usd_per_eur = df.year.map(lambda y: fx.usd_per_eur(int(y)))
    out = pd.DataFrame(index=df.index)
    out["agg_id"] = [stable_id(SOURCE, y, g, o, s) for y, g, o, s in
                     zip(df.year, df.geo_code, df.occupation_code, df.sex)]
    out["source"] = SOURCE
    out["source_dataset"] = DATASET
    out["license"] = LICENSE
    out["year"] = df.year

    out["occupation_scheme"] = "KldB2010"
    out["occupation_code"] = codes
    out["occupation_label"] = df.occupation_label.astype("string")
    out["occupation_level"] = codes.str.len().map(
        {2: "berufsgruppe", 3: "berufsuntergruppe", 4: "berufsgattung"}).fillna("total")
    out["normalized_job_code"] = mapped
    out["normalized_job_title"] = mapped.map(
        lambda c: role(c).label if pd.notna(c) else None)
    out["job_family"] = mapped.map(lambda c: role(c).family if pd.notna(c) else None)
    out["norm_method"] = method.fillna("unmatched")
    out["norm_confidence"] = mapped.notna().map({True: 0.8, False: 0.0})

    out["geo_level"] = df.geo_code.map(
        lambda g: "country" if g in {"D"} else ("region" if len(str(g)) <= 2 else "district"))
    out["geo_code"] = df.geo_code.astype("string")
    out["geo_name"] = df.geo_name.astype("string")
    out["country_iso2"] = "DE"
    out["region"] = df.geo_name.where(df.geo_code != "D")

    out["sex"] = df.sex
    out["worktime"] = "full_time"      # Kerngruppe = Vollzeit
    out["industry"] = "alle Wirtschaftszweige"
    out["employment_count"] = df.employment_count

    out["salary_currency"] = "EUR"
    out["salary_period"] = "month"
    out["p50_raw"] = df["median"]
    for q, name in ((0.10, "p10"), (0.25, "p25"), (0.75, "p75"), (0.90, "p90")):
        out[f"{name}_raw"] = quantiles[q]
    for stat in ("p10", "p25", "p50", "p75", "p90"):
        out[f"{stat}_annual_eur"] = out[f"{stat}_raw"] * MONTHS_PER_YEAR
    out["p50_annual_usd"] = out["p50_annual_eur"] * usd_per_eur
    out["quantile_method"] = "median_reported; p10/p25/p75/p90 interpolated_from_brackets"
    return schema.conform_aggregates(out)
